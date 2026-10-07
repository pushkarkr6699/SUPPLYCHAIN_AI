import io
import json
import zipfile
import numpy as np
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest
from config import ROOT
from services import upload_service as upload, visualization_service as viz


def csv_data():
    return b'Date,Segment,Amount,Quantity,Customer ID\n2026-01-01,Alpha,12,2,001\n2026-01-02,Beta,24,3,002\n'


def test_csv_preserves_identifiers_and_explicit_types():
    raw = upload.parse(csv_data(), 'data.csv')
    assert raw['Customer ID'].tolist() == ['001', '002']
    assert upload.numeric_candidates(raw) == ['Amount', 'Quantity']
    prepared = upload.prepare(raw, ['Amount','Quantity'], 'Date')
    assert prepared.Amount.sum() == 36 and prepared.Date.min() == pd.Timestamp('2026-01-01')
    assert raw.Amount.tolist() == ['12','24']
    for kind in viz.CHART_TYPES:
        groups=['Segment','Customer ID'] if kind in {'Grouped heatmap','Grouped bars','Stacked bars'} else ['Segment']
        metrics=['Amount','Quantity','Amount2'] if kind=='Bubble' else ['Amount','Quantity']
        data=prepared.assign(Amount2=[1.,2.])
        if kind=='Correlation heatmap':
            # Correlation requires three usable records.
            data=pd.concat([data,data],ignore_index=True)
        if kind in {'Moving average','Coordinate map','Ordered funnel'}:
            with pytest.raises(ValueError):viz.build(data,kind,groups,metrics)
        else:
            assert viz.build(data,kind,groups,metrics,'Sum' if kind=='Stacked bars' else 'Mean')['figure'].data


@pytest.mark.parametrize('content,name,options', [
    (b'a;b\n1;2\n','sample.csv',{'delimiter':';'}),
    (b'a\tb\n1\t2\n','sample.tsv',{}),
    ('a,b\nCaf\xe9,2\n'.encode('cp1252'),'sample.csv',{'encoding':'cp1252'}),
    (b'[{"a":"001","b":2},{"a":"002","b":null}]','sample.json',{}),
])
def test_supported_text_formats(content,name,options):
    result=upload.parse(content,name,**options)
    assert len(result)>0 and len(result.columns)==2


@pytest.mark.parametrize('content,name', [
    (b'', 'empty.csv'), (b'a,b\n','empty.csv'), (b'a,a\n1,2','duplicate.csv'),
    (b'a, a \n1,2','duplicate.csv'), (b'a,b\n1,2,3','fields.csv'),
    (b'\x00\x01', 'binary.csv'), (b'payload','model.pkl'),
    (b'[{"a":1,"a":2}]','duplicate.json'), (b'[{"a":[1]}]','nested.json'),
    (b'[{"a":NaN}]','nan.json'), (b'[{"a":1e400}]','inf.json'),
    (b'{"a":1}','object.json'), (b'not workbook','bad.xlsx'),
    (b'a\n'+b'x'*4097,'long.csv'), (b','.join([b'x'+str(i).encode() for i in range(101)])+b'\n','wide.csv'),
])
def test_invalid_files_fail_closed(content,name):
    with pytest.raises(ValueError):upload.parse(content,name)


def workbook(formula=False):
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet.title='Analysis'
    sheet.append(['Segment','Amount']);sheet.append(['Alpha','=1+1' if formula else 2])
    second=book.create_sheet('Another');second.append(['Metric']);second.append([4])
    out=io.BytesIO();book.save(out);return out.getvalue()


def test_workbook_sheet_selection_and_formulas():
    data=workbook()
    assert upload.sheets(data)==['Analysis','Another']
    assert upload.parse(data,'data.xlsx',sheet='Another').Metric.tolist()==[4]
    with pytest.raises(ValueError,match='Formula'):upload.parse(workbook(True),'formula.xlsx')
    with pytest.raises(ValueError):upload.parse(data,'data.xlsx',sheet='Missing')


def test_workbook_entities_and_expansion_are_rejected():
    for name,data in [('xl/worksheets/sheet1.xml',b'<!DOCTYPE x [<!ENTITY danger "x">]>'),
                      ('xl/worksheets/sheet1.xml',b'<c r="XFD999999999"/>'),
                      ('xl/vbaProject.bin',b'macro'),('xl/externalLinks/link.xml',b'link')]:
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as z:z.writestr(name,data)
        with pytest.raises(ValueError):upload.parse(out.getvalue(),'hostile.xlsx')
    with pytest.raises(ValueError):upload.parse(b'x'*(upload.MAX_BYTES+1),'large.csv')


def test_no_date_or_numeric_values_can_still_be_explored():
    prepared=upload.prepare(upload.parse(b'Segment\nAlpha\nBeta\nAlpha','raw.csv'),[])
    measure=prepared.attrs['generated_measure']
    table=viz.summary(prepared,['Segment'],[measure],operation='Sum')
    assert table[measure].sum()==3
    assert viz.build(prepared,'Donut',['Segment'],[measure])['figure'].data
    with pytest.raises(ValueError,match='date field'):viz.build(prepared,'Line',['Segment'],[measure])


def test_invalid_type_changes_do_not_fabricate_values():
    raw=upload.parse(b'Date,Metric\nnot-a-date,infinity\n','types.csv')
    with pytest.raises(ValueError):upload.prepare(raw,['Metric'])
    with pytest.raises(ValueError):upload.prepare(raw,[], 'Date')
    raw=upload.parse(b'Date,Metric\n01/02/2026,1\n','types.csv')
    assert upload.prepare(raw,['Metric'],'Date','Day first').Date.iloc[0]==pd.Timestamp('2026-02-01')


def source_for(model):
    if model=='delivery':
        from services.inference_service import input_rows
        return input_rows().head(3)
    if model=='demand':
        return pd.DataFrame({'DateOnly':pd.date_range('2026-01-01',periods=15).strftime('%Y-%m-%d'),
                             'Product':['Synthetic product']*15,'Category':['Synthetic category']*15,
                             'Department':['Synthetic department']*15,'Visits':np.arange(15,dtype=float)})
    path='data/cache/profitability_conversion_probe.csv' if model=='profitability' else 'data/cache/final_delivery_conversion_probe.csv'
    return pd.read_csv(ROOT/path).head(3).drop(columns='Expected Probability')


@pytest.mark.parametrize('model',list(upload.MODELS))
def test_actual_registered_models_score_mapped_uploaded_features(model):
    source=source_for(model)
    # Round-trip through the public data parser, which preserves CSV text.
    raw=upload.parse(source.to_csv(index=False).encode(), 'complete.csv')
    features=upload.model_contract(model)['features']
    renamed=raw.rename(columns={c:'Input '+str(i) for i,c in enumerate(features)})
    mapping={c:'Input '+str(i) for i,c in enumerate(features)}
    result=upload.predict(renamed,model,mapping)
    assert len(result)==(1 if model=='demand' else 3)
    assert result.attrs['uploaded'] and result.attrs['model_key']==model
    values=result['Predicted Visits'] if model=='demand' else result['Profitability Probability'] if model=='profitability' else result['Risk Probability']
    assert np.isfinite(values).all()
    if model=='demand':assert viz.units('Predicted Visits',result)=='web visits'
    if model!='demand':assert values.between(0,1).all()
    if model!='demand':
        numeric=upload.model_contract(model)['numeric']
        assert all(pd.api.types.is_numeric_dtype(result[mapping[c]]) for c in numeric)
        assert all(not pd.api.types.is_numeric_dtype(renamed[mapping[c]]) for c in numeric)
    assert list(pd.read_csv(io.BytesIO(upload.template(model))).columns)==features


@pytest.mark.parametrize('model',list(upload.MODELS))
def test_incomplete_features_and_duplicate_mapping_never_predict(model):
    raw=source_for(model);features=upload.model_contract(model)['features']
    with pytest.raises(ValueError,match='every required'):upload.predict(raw,model,{features[0]:features[0]})
    with pytest.raises(ValueError,match='distinct'):upload.map_features(raw,model,{c:features[0] for c in features})


def test_invalid_history_and_model_values_and_batch_limit():
    raw=source_for('demand');mapping={c:c for c in raw}
    with pytest.raises(ValueError,match='15 consecutive'):upload.predict(raw.head(14),'demand',mapping)
    raw.loc[0,'Visits']=-1
    with pytest.raises(ValueError):upload.predict(raw,'demand',mapping)
    raw=source_for('delivery');features=upload.model_contract('delivery')['features'];mapping={c:c for c in features}
    raw.loc[0,'Sales']=np.nan
    with pytest.raises(ValueError,match='complete finite'):upload.predict(raw,'delivery',mapping)
    raw.loc[0,'Sales']=10;raw.loc[0,'Order_Month']=14
    with pytest.raises(ValueError,match='Order_Month'):upload.predict(raw,'delivery',mapping)
    with pytest.raises(ValueError,match='5,000'):upload.predict(pd.concat([raw]*1700,ignore_index=True),'delivery',mapping)


def test_ai_anonymizes_headers_labels_and_requires_specific_consent(monkeypatch):
    frame=pd.DataFrame({'private@example.test':['Customer Alpha','Customer Beta'],'Secret revenue':[10.,20.]})
    payload,names=upload.ai_payload(frame,['private@example.test'],['Secret revenue'],'Mean','Explain the spread')
    wire=json.dumps(payload)
    for secret in ['private@example.test','Customer Alpha','Customer Beta','Secret revenue']:
        assert secret not in wire
    assert names['Secret revenue']=='Measure 1'
    calls=[]
    monkeypatch.setattr('services.ai_narration.narrate',lambda p:calls.append(p) or {'summary':'ok'})
    with pytest.raises(ValueError,match='Approve'):upload.narrate(frame,[],['Secret revenue'],'Mean','Explain',consent=False)
    assert not calls
    assert upload.narrate(frame,[],['Secret revenue'],'Mean','Explain',consent=True)['summary']=='ok'
    assert len(calls)==1


def test_formula_safe_exports():
    from services.export_service import csv_bytes
    frame=upload.parse(b'Name,Amount\n=HYPERLINK(""),1\n','raw.csv')
    assert "'=HYPERLINK" in csv_bytes(frame).decode('utf-8-sig')
    from services import export_service
    before=len(export_service._excel_cache)
    data=upload.excel_download(frame)
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(data))
    assert book.active['A2'].data_type=='s' and book.active['A2'].value.startswith("'=")
    assert len(export_service._excel_cache)==before


def test_row_limits_are_enforced_before_large_frames_are_created(monkeypatch):
    monkeypatch.setattr(upload,'MAX_ROWS',2)
    for data,name in [(b'Column\n1\n2\n3\n','rows.csv'),(b'[{"a":1},{"a":2},{"a":3}]','rows.json')]:
        with pytest.raises(ValueError):upload.parse(data,name)


def test_upload_route_renders_without_built_in_filters_and_explains_local_chat_consent(monkeypatch):
    from services import provider,sevika
    monkeypatch.setattr(provider,'SUPPLYCHAIN_PROVIDER','verified')
    app=AppTest.from_file(str(ROOT/'app.py'), default_timeout=30)
    for key,value in {'authenticated':True,'route':'uploads','route_initialized':True}.items():app.session_state[key]=value
    app.run()
    assert not app.exception
    assert any('Bring Your Data' in str(x.value) for x in app.get('html'))
    assert not any(x.key=='global_filters' for x in app.get('container'))
    context=sevika.build_context(pd.DataFrame(),'uploads','uploads',{},None)
    assert context['frame'].empty and 'file-specific consent checkbox' in context['guide']


def test_import_scope_revokes_consent_and_discards_stale_results():
    from views.uploads import _scope,clear
    app=AppTest.from_string("from views.uploads import _scope\nimport streamlit as st\n_scope('changed')\nst.write('ready')")
    for key,value in {'upload_scope':'old','upload_ai_consent':True,'upload_predictions':{'old':'scores'},'upload_raw':'private','viz_upload_raw_plan':{'old':'plan'}}.items():app.session_state[key]=value
    app.run();assert not app.exception
    for key in ['upload_ai_consent','upload_predictions','upload_raw','viz_upload_raw_plan']:
        assert key not in app.session_state


def test_upload_workflow_widgets_and_consent_are_isolated_per_file(monkeypatch):
    from types import SimpleNamespace
    import streamlit as st
    current=[csv_data()]
    monkeypatch.setattr(st,'file_uploader',lambda *a,**k:SimpleNamespace(name='sample.csv',getvalue=lambda:current[0]))
    monkeypatch.setattr('services.ai_narration.status',lambda:{'available':True})
    app=AppTest.from_string('from views.uploads import render\nrender()',default_timeout=30)
    app.run();assert not app.exception
    scope=app.session_state['upload_scope']
    app.checkbox(key='upload_ai_consent_'+scope).check().run()
    assert app.checkbox(key='upload_ai_consent_'+scope).value
    app.radio(key='upload_mode_'+scope).set_value('Trained-model predictions').run()
    assert app.selectbox(key='upload_model_'+scope)
    current[0]=b'Date,Segment,Amount,Quantity,Customer ID\n2026-02-01,New,99,5,004\n'
    app.run();assert not app.exception
    new_scope=app.session_state['upload_scope']
    assert new_scope!=scope
    assert app.radio(key='upload_mode_'+new_scope).value=='Raw data exploration'
    assert not app.checkbox(key='upload_ai_consent_'+new_scope).value
    assert 'upload_predictions' not in app.session_state
